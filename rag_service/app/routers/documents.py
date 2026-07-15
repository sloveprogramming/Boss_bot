"""RAG 服务路由 — 文档管理"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from typing import List

from app.database import get_db
from app.models.policy import PolicyDocument
from app.services.policy_rag import policy_rag_agent

router = APIRouter(prefix="/api/documents", tags=["文档管理"])


@router.post("/upload")
async def upload_document(
    title: str,
    category: str = "",
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """上传制度文件（支持 PDF/Word/TXT），自动解析并索引到 Milvus"""
    content = await file.read()
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""

    # 根据文件类型解析文本
    if ext == "txt":
        text = content.decode("utf-8", errors="replace")
    elif ext == "pdf":
        import fitz
        doc = fitz.open(stream=content, filetype="pdf")
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
    elif ext == "docx":
        import docx
        from io import BytesIO
        doc_obj = docx.Document(BytesIO(content))
        text = "\n".join(p.text for p in doc_obj.paragraphs)
    else:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {ext}")

    # 保存文档记录
    doc = PolicyDocument(
        title=title,
        category=category or "未分类",
        file_name=file.filename,
        content=text,
    )
    db.add(doc)
    await db.flush()

    # 自动索引到 Milvus
    try:
        await policy_rag_agent.index_document(doc.id, db)
    except Exception as e:
        doc.indexed = False
        await db.flush()
        return {
            "id": doc.id, "title": doc.title, "category": doc.category,
            "indexed": False, "warning": f"文件已保存，但索引失败: {str(e)}",
        }

    await db.refresh(doc)
    return {"id": doc.id, "title": doc.title, "category": doc.category, "indexed": True}


@router.get("")
async def list_documents(db: AsyncSession = Depends(get_db)):
    """获取所有制度文档列表"""
    result = await db.execute(
        select(PolicyDocument).order_by(PolicyDocument.created_at.desc())
    )
    docs = result.scalars().all()
    return [
        {
            "id": d.id, "title": d.title, "category": d.category,
            "file_name": d.file_name, "indexed": d.indexed,
            "chunk_count": d.chunk_count,
            "created_at": d.created_at.isoformat() if d.created_at else "",
        }
        for d in docs
    ]


@router.delete("/{doc_id}")
async def delete_document(doc_id: int, db: AsyncSession = Depends(get_db)):
    """删除制度文档"""
    result = await db.execute(
        select(PolicyDocument).where(PolicyDocument.id == doc_id)
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    await db.delete(doc)
    await db.flush()
    return {"status": "deleted"}


@router.post("/reindex")
async def reindex_all(db: AsyncSession = Depends(get_db)):
    """清空并重新索引所有文档"""
    result = await policy_rag_agent.reindex_all(db)
    return result
