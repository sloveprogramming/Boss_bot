"""简历解析与分析服务 — 文件解析 + LLM结构化提取 + 匹配度评估

依赖注入：
  __init__ 可接受 llm_client 参数用于注入/测试替换。
  未传入时回退到全局单例（兼容旧代码）。
"""
import json
import logging
import os
from typing import Optional
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.agents import RESUME_ANALYSIS_SYSTEM, RESUME_ANALYSIS_PROMPT
from app.models import Candidate, Resume

logger = logging.getLogger(__name__)


class ResumeAnalyzer:
    """简历分析器：负责文件解析和 AI 分析 — 支持依赖注入"""

    DEFAULT_UPLOAD_DIR = "uploads/resumes"  # 默认上传目录
    OCR_TEXT_THRESHOLD = 50

    def __init__(self, llm_client=None, upload_dir: Optional[str] = None):
        """构造简历分析器

        Args:
            llm_client: LLMClient 实例（可选，注入后用于测试替换）
            upload_dir: 上传目录路径（可选）
        """
        self._llm_client = llm_client
        self._upload_dir = upload_dir or self.DEFAULT_UPLOAD_DIR
        os.makedirs(self._upload_dir, exist_ok=True)
        self._ocr = None
        logger.info(f"简历分析器初始化完成，上传目录: {self._upload_dir}")

    @property
    def UPLOAD_DIR(self) -> str:
        return self._upload_dir

    def _get_llm(self):
        """获取 LLM 客户端 — 注入优先 → 回退全局单例"""
        if self._llm_client is not None:
            return self._llm_client
        from app.agents import llm_client
        return llm_client

    # ========== OCR 识别相关 ==========
    def _get_ocr(self):
        """惰性获取 PaddleOCR 引擎"""
        if self._ocr is None:
            logger.info("首次使用 PaddleOCR，正在加载模型...")
            from paddleocr import PaddleOCR
            self._ocr = PaddleOCR(use_angle_cls=True, lang='ch', use_gpu=False)
            logger.info("PaddleOCR 模型加载完成")
        return self._ocr

    def _ocr_image(self, image_bytes: bytes) -> str:
        """对图片进行 OCR 识别，返回纯文本"""
        from PIL import Image
        from io import BytesIO
        img = Image.open(BytesIO(image_bytes))
        logger.debug(f"OCR识别图片: 尺寸={img.size}, 模式={img.mode}")
        ocr = self._get_ocr()
        result = ocr.ocr(img, cls=True)
        if not result or not result[0]:
            logger.warning("OCR未识别到任何文字")
            return ""
        lines = []
        for block in result:
            for line in block:
                if line[1][1] > 0.3:
                    lines.append(line[1][0])
        text = "\n".join(lines)
        logger.info(f"OCR识别完成: 共 {len(lines)} 行文字, 总长度 {len(text)} 字符")
        return text

    # ========== 解析简历文件 ==========
    async def parse_resume_file(self, file_content: bytes, filename: str) -> str:
        """解析简历文件为纯文本"""
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        logger.info(f"开始解析简历文件: {filename} (类型: {ext}), 大小: {len(file_content)} 字节")

        if ext == "txt":
            text = file_content.decode("utf-8", errors="replace")
            logger.info(f"TXT解析完成: {len(text)} 字符")
            return text

        elif ext == "docx":
            import docx
            from io import BytesIO
            doc = docx.Document(BytesIO(file_content))
            text = "\n".join(p.text for p in doc.paragraphs)
            logger.info(f"DOCX解析完成: {len(doc.paragraphs)} 段落, {len(text)} 字符")
            return text

        elif ext == "pdf":
            import fitz
            doc = fitz.open(stream=file_content, filetype="pdf")
            page_count = len(doc)
            logger.info(f"PDF文档: {page_count} 页, 尝试提取文本...")

            text_parts = []
            for i, page in enumerate(doc):
                page_text = page.get_text()
                text_parts.append(page_text)
                logger.debug(f"  第{i+1}页: 提取到 {len(page_text)} 字符")
            text = "\n".join(text_parts)
            doc.close()

            if len(text.strip()) < self.OCR_TEXT_THRESHOLD:
                logger.warning(f"PDF文本提取仅 {len(text.strip())} 字符，判定为扫描件，切换OCR...")
                from PIL import Image
                doc = fitz.open(stream=file_content, filetype="pdf")
                ocr_pages = []
                for i, page in enumerate(doc):
                    logger.info(f"  OCR识别第 {i+1}/{page_count} 页...")
                    pix = page.get_pixmap(dpi=300)
                    img_bytes = pix.tobytes("png")
                    page_text = self._ocr_image(img_bytes)
                    ocr_pages.append(page_text)
                doc.close()
                text = "\n".join(ocr_pages)
                logger.info(f"PDF扫描件OCR完成: {len(text)} 字符")
            else:
                logger.info(f"PDF文本提取成功: {len(text)} 字符")
            return text

        elif ext in ("jpg", "jpeg", "png", "bmp", "tiff", "tif"):
            logger.info(f"图片文件OCR识别: {filename}")
            return self._ocr_image(file_content)

        else:
            logger.warning(f"未知文件类型: {ext}, 尝试作为纯文本处理")
            try:
                text = file_content.decode("utf-8", errors="replace")
                logger.info(f"作为纯文本解析成功: {len(text)} 字符")
                return text
            except Exception as e:
                logger.error(f"不支持的文件类型: {ext}, 错误: {e}")
                raise ValueError(f"不支持的文件类型: {ext}")

    # ========== 分析简历 ==========
    async def analyze_resume(self, resume_id: int, job_description: str,
                             db: AsyncSession) -> dict:
        """分析简历：调用 LLM 进行结构化提取和匹配度评估"""
        logger.info(f"开始AI分析简历 ID={resume_id}, 岗位描述={job_description[:50]}...")

        llm = self._get_llm()

        result = await db.execute(select(Resume).where(Resume.id == resume_id))
        resume = result.scalar_one_or_none()
        if not resume:
            logger.error(f"简历不存在: resume_id={resume_id}")
            raise ValueError(f"简历不存在: {resume_id}")
        if not resume.raw_text:
            logger.error(f"简历文本为空: resume_id={resume_id}")
            raise ValueError("简历文本为空，请先上传并解析")

        cand_result = await db.execute(
            select(Candidate).where(Candidate.id == resume.candidate_id)
        )
        candidate = cand_result.scalar_one_or_none()

        truncated_text = resume.raw_text[:15000]
        prompt = RESUME_ANALYSIS_PROMPT.format(
            resume_text=truncated_text,
            job_description=job_description or "通用岗位",
        )
        logger.info(f"调用LLM分析... 输入文本长度: {len(truncated_text)} 字符")

        response = await llm.chat(
            system=RESUME_ANALYSIS_SYSTEM,
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            temperature=0.2,
            max_tokens=4096,
        )

        try:
            result_data = json.loads(response)
        except json.JSONDecodeError:
            import re
            json_match = re.search(r'\{.*\}', response, re.DOTALL)
            if json_match:
                result_data = json.loads(json_match.group())
            else:
                raise ValueError(f"LLM返回格式异常: {response[:200]}...")

        analysis = result_data.get("analysis", {})
        match_analysis = result_data.get("match_analysis", {})
        summary = result_data.get("summary", "")

        resume.analysis_result = analysis
        resume.match_analysis = match_analysis
        resume.summary = summary

        if candidate:
            if match_analysis:
                candidate.match_score = match_analysis.get("overall_score")
            if analysis:
                basic_info = analysis.get("basic_info", {})
                if basic_info.get("name") and not candidate.name:
                    candidate.name = basic_info["name"]
                if basic_info.get("phone") and not candidate.phone:
                    candidate.phone = basic_info.get("phone")
                if basic_info.get("email") and not candidate.email:
                    candidate.email = basic_info.get("email")
                if analysis.get("skills"):
                    old_tags = list(candidate.tags or [])
                    candidate.tags = list(set(old_tags + analysis["skills"][:10]))
            if not candidate.status:
                candidate.status = "analyzed"

        import datetime as dt
        resume.analyzed_at = dt.datetime.utcnow()
        await db.flush()
        logger.info(f"简历分析全流程完成: resume_id={resume_id}")

        return {
            "analysis": analysis,
            "match_analysis": match_analysis,
            "summary": summary,
        }

    async def get_resume_analysis(self, candidate_id: int,
                                  db: AsyncSession) -> Optional[dict]:
        """获取候选人最新的简历分析结果"""
        result = await db.execute(
            select(Resume)
            .where(Resume.candidate_id == candidate_id)
            .order_by(Resume.created_at.desc())
            .limit(1)
        )
        resume = result.scalar_one_or_none()
        if not resume:
            return None

        return {
            "id": resume.id,
            "candidate_id": resume.candidate_id,
            "original_filename": resume.original_filename,
            "analysis": resume.analysis_result,
            "match_analysis": resume.match_analysis,
            "summary": resume.summary,
            "analyzed_at": resume.analyzed_at,
        }


# 全局单例 — 无参数构造，默认使用全局 llm_client（兼容旧代码）
resume_analyzer = ResumeAnalyzer()
