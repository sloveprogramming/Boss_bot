"""企业制度文档管理 API — 仅负责文档入库

RAG 问答、检索、聊天功能已迁移至 rag_service。
本模块仅保留向数据库存储文档内容的入口，
后续由 rag_service 负责索引到 Milvus 向量库。
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.database import get_db
from app.models.policy import PolicyDocument

router = APIRouter(prefix="/api/policy", tags=["制度文档入库"])


# ========== 文档入库 ==========

@router.post("/documents/upload")
async def upload_document(
    title: str,
    category: str = "",
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """上传制度文档并存储到数据库

    仅做文件解析 + 数据库入库。
    向量索引由 rag_service (端口 18003) 负责:
      POST http://localhost:18003/api/documents/reindex
    """
    content = await file.read()
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""

    # 根据文件类型解析文本
    if ext == "txt":
        text = content.decode("utf-8", errors="replace")
    elif ext == "pdf":
        import fitz
        doc = fitz.open(stream=content, filetype="pdf")
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
    elif ext == "docx":
        import docx
        from io import BytesIO
        doc_obj = docx.Document(BytesIO(content))
        text = "\n".join(p.text for p in doc_obj.paragraphs)
    else:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {ext}")

    # 保存文档记录（仅入库，不做向量索引）
    doc_record = PolicyDocument(
        title=title,
        category=category or "未分类",
        file_name=file.filename,
        content=text,
    )
    db.add(doc_record)
    await db.flush()
    await db.refresh(doc_record)

    return {
        "id": doc_record.id,
        "title": doc_record.title,
        "category": doc_record.category,
        "file_name": file.filename,
        "message": "文档已入库。如需索引到向量库，请调用 rag_service 的 /api/documents/reindex 接口。",
    }


# ========== 文档列表 ==========

@router.get("/documents")
async def list_documents(db: AsyncSession = Depends(get_db)):
    """获取所有已入库的制度文档列表"""
    result = await db.execute(
        select(PolicyDocument).order_by(PolicyDocument.created_at.desc())
    )
    docs = result.scalars().all()
    return [
        {
            "id": d.id,
            "title": d.title,
            "category": d.category,
            "file_name": d.file_name,
            "indexed": d.indexed,
            "chunk_count": d.chunk_count,
            "created_at": d.created_at.isoformat() if d.created_at else "",
        }
        for d in docs
    ]


@router.delete("/documents/{doc_id}")
async def delete_document(doc_id: int, db: AsyncSession = Depends(get_db)):
    """删除指定制度文档"""
    result = await db.execute(
        select(PolicyDocument).where(PolicyDocument.id == doc_id)
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    await db.delete(doc)
    await db.flush()
    return {"status": "deleted", "doc_id": doc_id}
